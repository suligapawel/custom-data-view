#!/usr/bin/env python3
"""
Wylicza odsetki dla obiektów wczytanych z pliku Excel i wpisuje je do tego samego pliku.

Przebieg:
  1. Wczytuje z pliku Excel obiekty (id, data do) ze wskazanych kolumn (adresy liter, np. A, CC).
     Wiersz jest pomijany, jeśli pasuje do któregokolwiek filtra wykluczającego KOLUMNA=WARTOŚĆ
     (np. A=NIE, BA=1, CC=TAK). Porównanie ignoruje wielkość liter i spacje na brzegach.
  2. Dla każdego obiektu pobiera dane finansowe:
       {FINANCE_BASE_URL}/latest/api/Finance/get/{id}?format=Tabular
     i uzupełnia pola:
       capital   <- $.LocalCurrencyData.Rows[DisplayName zaczyna się od 'Kapitał'].CurrentValue
       date_from <- $.LocalCurrencyData.Rows[DisplayName == 'Odsetki liczone od'].CurrentValue
  3. W tej samej iteracji pobiera wyliczone odsetki:
       {INTEREST_BASE_URL}/latest/api/interest?InterestValue={capital}&From={date_from}&To={date_to}&InterestType=2
     z nagłówkiem Authorization: Bearer <token> i z odpowiedzi bierze:
       interest  <- $.sumOfActiveInterest
  4. Wpisuje do pliku źródłowego, w wierszu z tym samym id, kapitał i datę "odsetki od dnia"
     (gdy udało się pobrać dane finansowe) oraz odsetki (gdy udało się je wyliczyć).
     Przed zapisem tworzy kopię zapasową pliku.

Przykład:
  python calculate_interest.py --input ~/dane/umowy.xlsx --id-column A --date-to-column CC \\
      --capital-column F --date-from-column G --interest-column E --exclude "A=NIE" --exclude "BA=1" --exclude "CC=TAK"

Filtry można też podać jednym parametrem, rozdzielone średnikiem: --exclude "A=NIE; BA=1; CC=TAK".
Zamiast '=' można użyć '->' (np. "A -> NIE"). Bez --exclude skrypt zapyta o filtry (Enter = brak).

Domyślnie dane zaczynają się od wiersza 2 (wiersz 1 to nagłówek); zmienia to --first-row.
Token Bearer jest zawsze pobierany interaktywnie przy uruchomieniu (wpisywanie niewidoczne).
Brakujące parametry (plik, kolumny) również są pobierane interaktywnie.
"""

import argparse
import getpass
import logging
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# --- Konfiguracja -----------------------------------------------------------

FINANCE_BASE_URL = "{PLACE_HOLDER_NA_URL}"
INTEREST_BASE_URL = "{PLACEHOLDER_NA_URL_2}"

FINANCE_URL_TEMPLATE = "{base}/latest/api/Finance/get/{id}?format=Tabular"
INTEREST_URL_TEMPLATE = "{base}/latest/api/interest"
INTEREST_TYPE = 2

CAPITAL_DISPLAY_NAME_PREFIX = "Kapitał"
DATE_FROM_DISPLAY_NAME = "Odsetki liczone od"
INTEREST_FIELD = "sumOfActiveInterest"

REQUEST_TIMEOUT_SECONDS = 30
DATE_FORMAT = "%Y-%m-%d"
EXCEL_DATE_FORMAT = "yyyy-mm-dd"

log = logging.getLogger("interest_calculator")


# --- Model ------------------------------------------------------------------

@dataclass
class FinanceItem:
    row_number: int
    id: str
    date_to: object
    date_from: object = None
    capital: object = None
    interest: object = None
    status: str = "PENDING"
    error: str = None


@dataclass
class Columns:
    id: str
    date_to: str
    capital: str
    date_from: str
    interest: str


@dataclass
class ExcludeFilter:
    column: str
    value: str

    def matches(self, row):
        return normalize_value(cell_value(row, column_index(self.column))) == self.value

    def __str__(self):
        return f"{self.column}={self.value}"


def parse_filters(texts):
    filters = []
    for text in texts:
        for part in text.split(";"):
            part = part.strip()
            if not part:
                continue
            separator = "->" if "->" in part else "="
            column, found, value = part.partition(separator)
            if not found or not column.strip():
                raise ValueError(f"Nieprawidłowy filtr: '{part}' (oczekiwano KOLUMNA=WARTOŚĆ, np. A=NIE)")
            column = column.strip().upper()
            column_index(column)
            filters.append(ExcludeFilter(column=column, value=normalize_value(value)))
    return filters


# --- Excel ------------------------------------------------------------------

def column_index(column_letter):
    letter = column_letter.strip().upper()
    try:
        return column_index_from_string(letter)
    except ValueError:
        raise ValueError(f"Nieprawidłowy adres kolumny: '{column_letter}' (oczekiwano np. A, CC)")


def normalize_id(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def normalize_value(value):
    if isinstance(value, bool):
        return "TAK" if value else "NIE"
    if isinstance(value, str):
        text = value.strip()
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            return text.upper()
        return normalize_id(number)
    return normalize_id(value).upper()


def cell_value(row, index):
    return row[index - 1] if index - 1 < len(row) else None


def load_items(path, sheet_name, columns, filters, first_row):
    id_index = column_index(columns.id)
    date_to_index = column_index(columns.date_to)

    log.info("Wczytywanie pliku: %s", path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        log.info("Arkusz: %s, kolumny: id=%s, data do=%s, kapitał=%s, odsetki od dnia=%s, odsetki=%s, od wiersza: %d",
                 sheet.title, columns.id, columns.date_to, columns.capital, columns.date_from, columns.interest,
                 first_row)
        log.info("Filtry wykluczające: %s", ", ".join(map(str, filters)) or "brak")

        items = []
        skipped_by_filter = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=first_row, values_only=True), start=first_row):
            item_id = normalize_id(cell_value(row, id_index))
            if not item_id:
                log.debug("Wiersz %d pominięty: brak id", row_number)
                continue
            matched = next((f for f in filters if f.matches(row)), None)
            if matched:
                log.debug("Wiersz %d (id=%s) pominięty: filtr %s", row_number, item_id, matched)
                skipped_by_filter += 1
                continue
            items.append(FinanceItem(row_number=row_number, id=item_id, date_to=cell_value(row, date_to_index)))
    finally:
        workbook.close()

    log.info("Wczytano %d obiektów do przetworzenia (pominięte przez filtry: %d)", len(items), skipped_by_filter)
    return items


def backup_file(path):
    backup_path = path.with_name(f"{path.stem}_kopia_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
    shutil.copy2(path, backup_path)
    log.info("Utworzono kopię zapasową: %s", backup_path)
    return backup_path


def write_results(path, sheet_name, columns, items):
    to_write = [item for item in items if item.capital is not None and item.date_from is not None]
    if not to_write:
        log.warning("Brak danych do zapisania - plik źródłowy nie został zmieniony")
        return

    backup_file(path)
    workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    id_index = column_index(columns.id)
    capital_index = column_index(columns.capital)
    date_from_index = column_index(columns.date_from)
    interest_index = column_index(columns.interest)

    written_finance = 0
    written_interest = 0
    for item in to_write:
        row_id = normalize_id(sheet.cell(row=item.row_number, column=id_index).value)
        if row_id != item.id:
            item.status = "BŁĄD"
            item.error = f"Id w wierszu {item.row_number} ('{row_id}') nie zgadza się z obiektem"
            log.error("id=%s - %s, dane nie zostały wpisane", item.id, item.error)
            continue
        sheet.cell(row=item.row_number, column=capital_index).value = item.capital
        date_cell = sheet.cell(row=item.row_number, column=date_from_index)
        date_cell.value = datetime.strptime(item.date_from, DATE_FORMAT).date()
        date_cell.number_format = EXCEL_DATE_FORMAT
        written_finance += 1
        if item.status == "OK":
            sheet.cell(row=item.row_number, column=interest_index).value = item.interest
            written_interest += 1

    try:
        workbook.save(path)
    except PermissionError:
        log.error("Nie można zapisać pliku %s - zamknij go w Excelu i uruchom skrypt ponownie", path)
        raise
    log.info("Wpisano kapitał i odsetki od dnia dla %d wierszy (kolumny %s, %s) oraz odsetki dla %d wierszy "
             "(kolumna %s) w pliku: %s",
             written_finance, columns.capital, columns.date_from, written_interest, columns.interest, path)


# --- Konwersje --------------------------------------------------------------

def format_date(value):
    if value is None or str(value).strip() == "":
        raise ValueError("Brak daty")
    if isinstance(value, (datetime, date)):
        return value.strftime(DATE_FORMAT)
    text = str(value).strip()
    for pattern in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y.%m.%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:19] if "T" in text else text, pattern).strftime(DATE_FORMAT)
        except ValueError:
            continue
    raise ValueError(f"Nieobsługiwany format daty: '{text}'")


def parse_amount(value):
    if value is None:
        raise ValueError("Brak kwoty")
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(" ", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"Nieprawidłowa kwota: '{value}'")


# --- REST API ---------------------------------------------------------------

def find_row_value(payload, display_name, prefix=False):
    rows = (payload.get("LocalCurrencyData") or {}).get("Rows") or []
    matches = []
    for row in rows:
        name = str(row.get("DisplayName", "")).strip()
        if name.startswith(display_name) if prefix else name == display_name:
            matches.append(row)
    if not matches:
        description = f"zaczynającego się od '{display_name}'" if prefix else f"'{display_name}'"
        raise ValueError(f"Brak wiersza {description} w LocalCurrencyData.Rows")
    if len(matches) > 1:
        log.warning("Kilka wierszy pasuje do '%s' (%s) - użyto pierwszego: '%s'", display_name,
                    ", ".join(str(row.get("DisplayName")) for row in matches), matches[0].get("DisplayName"))
    return matches[0].get("CurrentValue")


def fetch_finance(session, item):
    url = FINANCE_URL_TEMPLATE.format(base=FINANCE_BASE_URL.rstrip("/"), id=requests.utils.quote(item.id, safe=""))
    response = session.get(url, timeout=REQUEST_TIMEOUT_SECONDS)
    response.raise_for_status()
    payload = response.json()
    item.capital = parse_amount(find_row_value(payload, CAPITAL_DISPLAY_NAME_PREFIX, prefix=True))
    item.date_from = format_date(find_row_value(payload, DATE_FROM_DISPLAY_NAME))


def fetch_interest(session, item, token):
    url = INTEREST_URL_TEMPLATE.format(base=INTEREST_BASE_URL.rstrip("/"))
    params = {
        "InterestValue": item.capital,
        "From": item.date_from,
        "To": format_date(item.date_to),
        "InterestType": INTEREST_TYPE,
    }
    response = session.get(
        url,
        params=params,
        headers={"Authorization": f"Bearer {token}"},
        timeout=REQUEST_TIMEOUT_SECONDS,
    )
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, dict) or INTEREST_FIELD not in payload:
        raise ValueError(f"Brak pola '{INTEREST_FIELD}' w odpowiedzi API odsetek")
    item.interest = payload[INTEREST_FIELD]


def describe_error(error):
    if isinstance(error, requests.HTTPError) and error.response is not None:
        return f"HTTP {error.response.status_code}"
    if isinstance(error, requests.Timeout):
        return "Przekroczono limit czasu żądania"
    if isinstance(error, requests.ConnectionError):
        return "Błąd połączenia"
    if isinstance(error, requests.RequestException):
        return type(error).__name__
    return str(error)


def process_items(items, token):
    total = len(items)
    width = len(str(total))
    with requests.Session() as session:
        session.headers["Accept"] = "application/json"
        for index, item in enumerate(items, start=1):
            prefix = f"[{index:>{width}}/{total}] ({index * 100 // total:>3}%) wiersz {item.row_number}, id={item.id}"
            try:
                log.info("%s - pobieranie danych finansowych", prefix)
                fetch_finance(session, item)
                log.info("%s - kapitał=%s, data od=%s, data do=%s",
                         prefix, item.capital, item.date_from, format_date(item.date_to))

                log.info("%s - pobieranie odsetek", prefix)
                fetch_interest(session, item, token)
                item.status = "OK"
                log.info("%s - odsetki=%s", prefix, item.interest)
            except Exception as error:
                item.status = "BŁĄD"
                item.error = describe_error(error)
                log.error("%s - %s", prefix, item.error)


def log_summary(items, elapsed):
    failed = [item for item in items if item.status != "OK"]
    log.info("Zakończono: %d obiektów, OK: %d, błędy/nieprzetworzone: %d, czas: %.1fs",
             len(items), len(items) - len(failed), len(failed), elapsed)
    for item in failed:
        log.warning("Bez odsetek: wiersz %d, id=%s - %s", item.row_number, item.id, item.error or "nieprzetworzony")


# --- CLI --------------------------------------------------------------------

def ask(value, prompt):
    while not value:
        value = input(f"{prompt}: ").strip()
    return value


def parse_args():
    parser = argparse.ArgumentParser(description="Wylicza odsetki dla obiektów z pliku Excel i wpisuje je do tego pliku.")
    parser.add_argument("--input", help="Ścieżka do pliku Excel wraz z nazwą pliku")
    parser.add_argument("--sheet", help="Nazwa arkusza (domyślnie aktywny)")
    parser.add_argument("--id-column", help="Adres kolumny z id, np. A")
    parser.add_argument("--date-to-column", help="Adres kolumny z datą do, np. CC")
    parser.add_argument("--exclude", action="append", default=[],
                        help="Filtr wykluczający KOLUMNA=WARTOŚĆ, np. A=NIE; można powtarzać lub rozdzielić ';'")
    parser.add_argument("--capital-column", help="Adres kolumny, do której zostanie wpisany kapitał, np. F")
    parser.add_argument("--date-from-column", help="Adres kolumny, do której zostanie wpisana data odsetek od dnia, np. G")
    parser.add_argument("--interest-column", help="Adres kolumny, do której zostaną wpisane odsetki, np. E")
    parser.add_argument("--first-row", type=int, default=2, help="Pierwszy wiersz z danymi (domyślnie 2)")
    parser.add_argument("--log-file", help="Plik logu (domyślnie <wejście>_odsetki.log)")
    parser.add_argument("--verbose", action="store_true", help="Szczegółowe logowanie")
    return parser.parse_args()


def configure_logging(log_file, verbose):
    formatter = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%Y-%m-%d %H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(formatter)
    file_handler = logging.FileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(formatter)
    log.setLevel(logging.DEBUG if verbose else logging.INFO)
    log.addHandler(console)
    log.addHandler(file_handler)


def main():
    args = parse_args()
    input_path = Path(ask(args.input, "Ścieżka do pliku Excel")).expanduser()
    columns = Columns(
        id=ask(args.id_column, "Adres kolumny z id (np. A)").upper(),
        date_to=ask(args.date_to_column, "Adres kolumny z datą do (np. CC)").upper(),
        capital=ask(args.capital_column, "Adres kolumny na kapitał (np. F)").upper(),
        date_from=ask(args.date_from_column, "Adres kolumny na odsetki od dnia (np. G)").upper(),
        interest=ask(args.interest_column, "Adres kolumny na odsetki (np. E)").upper(),
    )
    filter_texts = args.exclude or [input("Filtry wykluczające, np. A=NIE; BA=1; CC=TAK (Enter = brak): ")]
    try:
        filters = parse_filters(filter_texts)
    except ValueError as error:
        print(error, file=sys.stderr)
        return 2
    token = getpass.getpass("Token Bearer (niewidoczny): ").strip()
    if not token:
        print("Token jest wymagany.", file=sys.stderr)
        return 2

    log_path = Path(args.log_file).expanduser() if args.log_file else input_path.with_name(f"{input_path.stem}_odsetki.log")
    configure_logging(log_path, args.verbose)
    log.info("Log zapisywany do: %s", log_path)

    if not input_path.is_file():
        log.error("Plik nie istnieje: %s", input_path)
        return 1

    try:
        for column in (columns.capital, columns.date_from, columns.interest):
            column_index(column)
        items = load_items(input_path, args.sheet, columns, filters, args.first_row)
    except Exception as error:
        log.error("Nie udało się wczytać pliku: %s", error)
        return 1

    if not items:
        log.warning("Brak obiektów do przetworzenia")
        return 0

    started = time.monotonic()
    try:
        process_items(items, token)
    except KeyboardInterrupt:
        log.warning("Przerwano przez użytkownika - zapisuję dane pobrane do tej pory")
    finally:
        try:
            write_results(input_path, args.sheet, columns, items)
        except Exception as error:
            log.error("Nie udało się zapisać odsetek do pliku: %s", type(error).__name__)
            return 1
        finally:
            log_summary(items, time.monotonic() - started)
    return 0


if __name__ == "__main__":
    sys.exit(main())
